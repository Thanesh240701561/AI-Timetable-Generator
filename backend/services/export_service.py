import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExportService:
    @staticmethod
    def export_timetable_excel(timetable_data, export_type="Complete Timetable", conflicts=None):
        """
        Generates an in-memory Excel file containing formatted sheets.
        Returns bytes or BytesIO stream.
        """
        wb = openpyxl.Workbook()
        wb.remove(wb.active) # Remove default sheet

        df = pd.DataFrame(timetable_data)
        if df.empty:
            ws = wb.create_sheet("Timetable")
            ws.append(["No timetable data available."])
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output

        header_fill = PatternFill(start_color="171F35", end_color="171F35", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        class_fill = PatternFill(start_color="202A44", end_color="202A44", fill_type="solid")
        break_fill = PatternFill(start_color="11182B", end_color="11182B", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='2D3955'),
            right=Side(style='thin', color='2D3955'),
            top=Side(style='thin', color='2D3955'),
            bottom=Side(style='thin', color='2D3955')
        )

        def style_worksheet(ws):
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if str(cell.value) == "Class":
                        cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
                        cell.font = Font(color="FFFFFF", bold=True)
                    elif str(cell.value) == "Break":
                        cell.fill = PatternFill(start_color="7C5CFC", end_color="7C5CFC", fill_type="solid")
                        cell.font = Font(color="FFFFFF", bold=True)

        DAY_ORDER = {
            "Monday": 1,
            "Tuesday": 2,
            "Wednesday": 3,
            "Thursday": 4,
            "Friday": 5,
            "Saturday": 6,
            "Sunday": 7
        }
        df['day_num'] = df['day'].map(lambda d: DAY_ORDER.get(str(d).strip(), 99))
        df['period_num'] = pd.to_numeric(df['period'], errors='coerce').fillna(99)

        # 1. Complete Timetable Sheet
        ws_complete = wb.create_sheet("Complete Timetable")
        cols = ["timetable_id", "day", "period", "start_time", "end_time", "section", "group_id", "subject_code", "subject_name", "faculty_name", "room_name", "status"]
        ws_complete.append([c.replace("_", " ").title() for c in cols])
        df_complete = df.sort_values(by=['day_num', 'period_num', 'start_time'])
        for _, row in df_complete.iterrows():
            ws_complete.append([row.get(c, "") for c in cols])
        style_worksheet(ws_complete)

        # 2. Group-Wise Sheet
        ws_group = wb.create_sheet("Group Timetables")
        ws_group.append([c.replace("_", " ").title() for c in cols])
        df_group = df.sort_values(by=['group_id', 'day_num', 'period_num', 'start_time'])
        for _, row in df_group.iterrows():
            ws_group.append([row.get(c, "") for c in cols])
        style_worksheet(ws_group)

        # 3. Conflicts Sheet (if any)
        if conflicts:
            ws_conflicts = wb.create_sheet("Conflict Report")
            c_cols = ["id", "type", "issue", "subject", "faculty", "room", "group", "day", "period", "action"]
            ws_conflicts.append([c.replace("_", " ").title() for c in c_cols])
            for c_item in conflicts:
                ws_conflicts.append([c_item.get(c, "") for c in c_cols])
            style_worksheet(ws_conflicts)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
