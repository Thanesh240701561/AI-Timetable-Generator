import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REQUIRED_SHEETS = [
    "Students", "Subjects", "Faculty", "TeacherGroups", "Rooms", "TimeSlots", "Availability"
]

SHEET_COLUMNS = {
    "Students": ["student_id", "student_name", "section", "group_id"],
    "Subjects": ["subject_id", "subject_code", "subject_name", "periods_per_week", "subject_type"],
    "Faculty": ["faculty_id", "faculty_name", "subject_id", "group_id"],
    "TeacherGroups": ["group_id", "subject_id", "faculty_id", "room_id", "max_capacity"],
    "Rooms": ["room_id", "room_name", "room_type", "capacity"],
    "TimeSlots": ["slot_id", "day", "period", "start_time", "end_time"],
    "Availability": ["faculty_id", "day", "period", "available"]
}

class ExcelService:
    def __init__(self, file_path=None):
        self.file_path = file_path

    def load_data(self, file_path=None):
        target_path = file_path or self.file_path
        if not target_path or not os.path.exists(target_path):
            raise FileNotFoundError("Excel file not found.")

        excel_data = {}
        xls = pd.ExcelFile(target_path)
        for sheet in REQUIRED_SHEETS:
            if sheet in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet)
                df.columns = [str(col).strip() for col in df.columns]
                for col in df.columns:
                    if df[col].dtype == object:
                        df[col] = df[col].astype(str).str.strip()
                excel_data[sheet] = df
            else:
                excel_data[sheet] = pd.DataFrame()
        return excel_data

    def validate_excel(self, file_path=None):
        target_path = file_path or self.file_path
        errors = []
        warnings = []

        if not target_path or not os.path.exists(target_path):
            return {"valid": False, "errors": ["Excel file does not exist."], "warnings": []}

        try:
            xls = pd.ExcelFile(target_path)
        except Exception as e:
            return {"valid": False, "errors": [f"Failed to read Excel file: {str(e)}"], "warnings": []}

        missing_sheets = [s for s in REQUIRED_SHEETS if s not in xls.sheet_names]
        if missing_sheets:
            errors.append(f"Missing required sheets: {', '.join(missing_sheets)}")
            return {"valid": False, "errors": errors, "warnings": warnings}

        data = {}
        for sheet in REQUIRED_SHEETS:
            df = pd.read_excel(xls, sheet_name=sheet)
            df.columns = [str(col).strip() for col in df.columns]
            data[sheet] = df

            required_cols = SHEET_COLUMNS[sheet]
            missing_cols = [c for c in required_cols if c not in df.columns]
            if missing_cols:
                errors.append(f"Sheet '{sheet}' is missing columns: {', '.join(missing_cols)}")
            elif df.empty:
                warnings.append(f"Sheet '{sheet}' is empty.")

        if errors:
            return {"valid": False, "errors": errors, "warnings": warnings}

        students_df = data["Students"]
        subjects_df = data["Subjects"]
        faculty_df = data["Faculty"]
        groups_df = data["TeacherGroups"]
        rooms_df = data["Rooms"]

        if not students_df.empty and "student_id" in students_df.columns:
            dups = students_df[students_df.duplicated("student_id")]["student_id"].tolist()
            if dups:
                errors.append(f"Duplicate student_ids found: {', '.join(map(str, dups))}")

        if not faculty_df.empty and "faculty_id" in faculty_df.columns:
            dups = faculty_df[faculty_df.duplicated("faculty_id")]["faculty_id"].tolist()
            if dups:
                errors.append(f"Duplicate faculty_ids found: {', '.join(map(str, dups))}")

        if not rooms_df.empty and "room_id" in rooms_df.columns:
            dups = rooms_df[rooms_df.duplicated("room_id")]["room_id"].tolist()
            if dups:
                errors.append(f"Duplicate room_ids found: {', '.join(map(str, dups))}")

        valid_subject_ids = set(subjects_df["subject_id"].dropna()) if "subject_id" in subjects_df.columns else set()
        valid_group_ids = set(groups_df["group_id"].dropna()) if "group_id" in groups_df.columns else set()
        valid_room_ids = set(rooms_df["room_id"].dropna()) if "room_id" in rooms_df.columns else set()

        if "subject_id" in faculty_df.columns:
            invalid_sub = faculty_df[~faculty_df["subject_id"].isin(valid_subject_ids)]["subject_id"].tolist()
            if invalid_sub:
                errors.append(f"Faculty sheet has invalid subject_ids: {', '.join(set(map(str, invalid_sub)))}")

        if "group_id" in students_df.columns:
            invalid_grp = students_df[~students_df["group_id"].isin(valid_group_ids)]["group_id"].tolist()
            if invalid_grp:
                errors.append(f"Students sheet has invalid group_ids: {', '.join(set(map(str, invalid_grp)))}")

        if "room_id" in groups_df.columns:
            invalid_rm = groups_df[~groups_df["room_id"].isin(valid_room_ids)]["room_id"].tolist()
            if invalid_rm:
                errors.append(f"TeacherGroups sheet has invalid room_ids: {', '.join(set(map(str, invalid_rm)))}")

        is_valid = len(errors) == 0
        return {"valid": is_valid, "errors": errors, "warnings": warnings}

    @staticmethod
    def generate_sample_workbook(save_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Students (60 Student Records across 5 Sections A, B, C, D, E mapped to Groups G1..G5)
        ws_students = wb.create_sheet("Students")
        ws_students.append(["student_id", "student_name", "section", "group_id"])
        
        students_data = []
        names = [
            "Rahul Sharma", "Priya Singh", "Aman Verma", "Neha Patel", "Rohan Kumar", "Ananya Roy",
            "Karan Gupta", "Sneha Rao", "Vikram Malhotra", "Pooja Hegde", "Aditya Joshi", "Riya Sen",
            "Siddharth Jain", "Kavya Nair", "Varun Dhawan", "Shraddha Kapoor", "Ayushmann Khurrana", "Kriti Sanon",
            "Kartik Aaryan", "Kiara Advani", "Ranbir Kapoor", "Alia Bhatt", "Ranveer Singh", "Deepika Padukone",
            "Vicky Kaushal", "Katrina Kaif", "Shahid Kapoor", "Mira Rajput", "Rajkummar Rao", "Patralekhaa",
            "Nawazuddin Siddiqui", "Pankaj Tripathi", "Manoj Bajpayee", "Jaideep Ahlawat", "Vijay Varma", "Radhika Apte",
            "Ayush Sharma", "Sanya Malhotra", "Fatima Sana", "Taapsee Pannu", "Yami Gautam", "Bhumi Pednekar",
            "Vidyut Jammwal", "Tiger Shroff", "Disha Patani", "Hrithik Roshan", "John Abraham", "Jacqueline Fernandez",
            "Abhishek Bachchan", "Aishwarya Rai", "Suryakumar Yadav", "Hardik Pandya", "Jasprit Bumrah", "Rohit Sharma",
            "Virat Kohli", "KL Rahul", "Rishabh Pant", "Shreyas Iyer", "Mohammed Shami", "Ravindra Jadeja"
        ]
        
        sections = ["Section A", "Section B", "Section C", "Section D", "Section E"]
        groups = ["G1", "G2", "G3", "G4", "G5"]

        for i in range(60):
            stu_id = f"STU{i+1:03d}"
            stu_name = names[i % len(names)]
            sec_idx = i // 12
            section = sections[sec_idx]
            group_id = groups[sec_idx]
            students_data.append([stu_id, stu_name, section, group_id])

        for row in students_data:
            ws_students.append(row)

        # 2. Subjects (EXACTLY 6 SUBJECTS, each 4 periods/week)
        ws_subjects = wb.create_sheet("Subjects")
        ws_subjects.append(["subject_id", "subject_code", "subject_name", "periods_per_week", "subject_type"])
        subjects_data = [
            ["SUB101", "CS23531", "Data Structures & Algorithms", 4, "Lecture"],
            ["SUB102", "CS23532", "Database Management Systems", 4, "Lecture"],
            ["SUB103", "CS23533", "Operating Systems", 4, "Lecture"],
            ["SUB104", "CS23512", "Artificial Intelligence", 4, "Lecture"],
            ["SUB105", "CS23514", "Computer Networks", 4, "Lecture"],
            ["SUB106", "CS23516", "Discrete Mathematics", 4, "Lecture"]
        ]
        for row in subjects_data:
            ws_subjects.append(row)

        # 3. Faculty (EXACTLY 5 TEACHERS FOR EACH SUBJECT = 30 Faculty Members Total!)
        ws_faculty = wb.create_sheet("Faculty")
        ws_faculty.append(["faculty_id", "faculty_name", "subject_id", "group_id"])
        
        faculty_data = []
        subject_teachers = {
            "SUB101": [("FAC001", "Dr. Smith (DS Group A)", "G1"), ("FAC002", "Dr. Smith (DS Group B)", "G2"), ("FAC003", "Dr. Smith (DS Group C)", "G3"), ("FAC004", "Dr. Smith (DS Group D)", "G4"), ("FAC005", "Dr. Smith (DS Group E)", "G5")],
            "SUB102": [("FAC006", "Dr. Kumar (DBMS Group A)", "G1"), ("FAC007", "Dr. Kumar (DBMS Group B)", "G2"), ("FAC008", "Dr. Kumar (DBMS Group C)", "G3"), ("FAC009", "Dr. Kumar (DBMS Group D)", "G4"), ("FAC010", "Dr. Kumar (DBMS Group E)", "G5")],
            "SUB103": [("FAC011", "Dr. Lee (OS Group A)", "G1"), ("FAC012", "Dr. Lee (OS Group B)", "G2"), ("FAC013", "Dr. Lee (OS Group C)", "G3"), ("FAC014", "Dr. Lee (OS Group D)", "G4"), ("FAC015", "Dr. Lee (OS Group E)", "G5")],
            "SUB104": [("FAC016", "Dr. Patel (AI Group A)", "G1"), ("FAC017", "Dr. Patel (AI Group B)", "G2"), ("FAC018", "Dr. Patel (AI Group C)", "G3"), ("FAC019", "Dr. Patel (AI Group D)", "G4"), ("FAC020", "Dr. Patel (AI Group E)", "G5")],
            "SUB105": [("FAC021", "Dr. Wilson (CN Group A)", "G1"), ("FAC022", "Dr. Wilson (CN Group B)", "G2"), ("FAC023", "Dr. Wilson (CN Group C)", "G3"), ("FAC024", "Dr. Wilson (CN Group D)", "G4"), ("FAC025", "Dr. Wilson (CN Group E)", "G5")],
            "SUB106": [("FAC026", "Dr. Davis (DM Group A)", "G1"), ("FAC027", "Dr. Davis (DM Group B)", "G2"), ("FAC028", "Dr. Davis (DM Group C)", "G3"), ("FAC029", "Dr. Davis (DM Group D)", "G4"), ("FAC030", "Dr. Davis (DM Group E)", "G5")]
        }

        for sub_id, teachers in subject_teachers.items():
            for fac_id, fac_name, grp_id in teachers:
                faculty_data.append([fac_id, fac_name, sub_id, grp_id])

        for row in faculty_data:
            ws_faculty.append(row)

        # 4. TeacherGroups (5 Groups x 6 Subjects = 30 Group Mappings)
        ws_groups = wb.create_sheet("TeacherGroups")
        ws_groups.append(["group_id", "subject_id", "faculty_id", "room_id", "max_capacity"])
        
        groups_data = []
        room_list = ["R101", "R102", "R103", "R104", "R105", "R106", "R107", "R108", "R109", "R110"]
        
        for grp_idx, grp_id in enumerate(["G1", "G2", "G3", "G4", "G5"]):
            for sub_idx, sub_id in enumerate(["SUB101", "SUB102", "SUB103", "SUB104", "SUB105", "SUB106"]):
                fac_info = subject_teachers[sub_id][grp_idx]
                fac_id = fac_info[0]
                room_id = room_list[(grp_idx + sub_idx) % len(room_list)]
                groups_data.append([grp_id, sub_id, fac_id, room_id, 60])

        for row in groups_data:
            ws_groups.append(row)

        # 5. Rooms (10 Lecture Halls)
        ws_rooms = wb.create_sheet("Rooms")
        ws_rooms.append(["room_id", "room_name", "room_type", "capacity"])
        rooms_data = [
            ["R101", "Hall A101", "Lecture", 60],
            ["R102", "Hall A102", "Lecture", 60],
            ["R103", "Hall A201", "Lecture", 60],
            ["R104", "Hall A202", "Lecture", 60],
            ["R105", "Hall B101", "Lecture", 60],
            ["R106", "Hall B102", "Lecture", 60],
            ["R107", "Hall B201", "Lecture", 60],
            ["R108", "Hall B202", "Lecture", 60],
            ["R109", "Hall C101", "Lecture", 60],
            ["R110", "Hall C102", "Lecture", 60]
        ]
        for row in rooms_data:
            ws_rooms.append(row)

        # 6. TimeSlots (30 Time Slots: 5 Days x 6 Periods)
        ws_slots = wb.create_sheet("TimeSlots")
        ws_slots.append(["slot_id", "day", "period", "start_time", "end_time"])
        days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
        periods = [
            (1, "08:00", "08:50"),
            (2, "09:00", "09:50"),
            (3, "10:00", "10:50"),
            (4, "11:00", "11:50"),
            (5, "13:00", "13:50"),
            (6, "14:00", "14:50")
        ]
        slot_id = 1
        for day in days:
            for p_num, s_time, e_time in periods:
                ws_slots.append([f"TS{slot_id:03d}", day, p_num, s_time, e_time])
                slot_id += 1

        # 7. Availability Matrix (All 30 Faculty Members)
        ws_avail = wb.create_sheet("Availability")
        ws_avail.append(["faculty_id", "day", "period", "available"])
        for fac_id in [f"FAC{i+1:03d}" for i in range(30)]:
            for day in days:
                for p_num in range(1, 7):
                    avail = 1
                    if fac_id == "FAC001" and day == "Friday" and p_num == 6:
                        avail = 0
                    ws_avail.append([fac_id, day, p_num, avail])

        # Style header rows
        header_fill = PatternFill(start_color="171F35", end_color="171F35", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        for sheet in wb.worksheets:
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

        dirname = os.path.dirname(save_path)
        if dirname:
            os.makedirs(dirname, exist_ok=True)
        wb.save(save_path)
        return save_path
